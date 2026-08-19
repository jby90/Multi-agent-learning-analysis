from __future__ import annotations

from collections import Counter, defaultdict

import pytest

from eval.v4_cases import load_gold_standard
from eval.v4_recompute import normalize_v4_case, normalize_v4_gold
from eval.v4_workbook import export_v4_workbook, read_human_review_workbook


def _diagnosis(point: str, difficulty: str, msg_id: str) -> dict:
    return {
        "msg_id": msg_id,
        "step": 4,
        "role": "produce",
        "agent": "diagnosis",
        "payload": {
            "type": "profile_assessment",
            "content": {
                "event": "diagnosis_ready",
                "selected_knowledge_point": point,
                "selected_difficulty": difficulty,
                "knowledge_point_plan": [{
                    "knowledge_point": point,
                    "initial_difficulty": difficulty,
                    "route_reason": "冻结前测与探针的确定性路由结果",
                    "evidence_ids": ["PT-1"],
                }],
            },
        },
    }


def _case() -> dict:
    return {
        "run_id": "V4-A",
        "seed_id": "seed_A",
        "case_id": "E2E-011",
        "route_mode": "production",
        "profile_id": "craft_engineer",
        "target_knowledge_point": "完成率计算",
        "sessions": [
            {
                "session_id": "pre-1",
                "knowledge_point": "计划量与实际量口径",
                "messages": [_diagnosis("计划量与实际量口径", "applied", "m-1")],
            },
            {
                "session_id": "target-1",
                "knowledge_point": "完成率计算",
                "messages": [_diagnosis("完成率计算", "basic", "m-2")],
            },
        ],
    }


def test_v4_normalizer_keeps_all_sessions_for_facts_but_scores_the_target_session():
    normalized = normalize_v4_case(_case())

    assert [row["session_id"] for row in normalized.fact_runs] == ["pre-1", "target-1"]
    assert normalized.target_run["session_id"] == "target-1"
    assert normalized.target_run["messages"][0]["payload"]["content"][
        "selected_knowledge_point"
    ] == "完成率计算"


def test_v4_normalizer_refuses_a_missing_target_session():
    case = _case()
    case["target_knowledge_point"] = "不存在的目标"

    with pytest.raises(ValueError, match="target session"):
        normalize_v4_case(case)


def test_v4_gold_designates_exactly_ten_by_three_real_coverage_cells():
    rows = normalize_v4_gold(load_gold_standard())
    coverage = [row for row in rows if row["计入覆盖率"] == "是"]

    assert len(coverage) == 30
    by_point: dict[str, set[str]] = defaultdict(set)
    for row in coverage:
        by_point[row["目标知识点"]].add(row["覆盖格"].rsplit("-", 1)[-1])
    assert len(by_point) == 10
    assert all(values == {"BASIC", "APPLIED", "ADVANCED"} for values in by_point.values())
    assert Counter(row["learner_script_id"] for row in coverage) == {
        "S-STEP-UP-B": 10,
        "S-STEP-UP-A": 10,
        "S-REBUTTAL": 10,
    }
    assert sum(row["预期适配节点数"] for row in rows) == 130
    assert {
        row["预期适配节点数"]
        for row in rows
        if row["learner_script_id"] == "S-REFRESH"
    } == {2}


def test_v4_workbook_keeps_auto_results_and_reads_only_human_columns(tmp_path):
    report = {
        "mode": "AUTO_PRELIMINARY",
        "combined": {
            "label_status": "AUTOMATIC_INITIAL_ESTIMATE_NOT_FINAL",
            "metrics": {
                key: {"numerator": 1, "denominator": 2, "percentage": 50.0, "denominator_zero": False}
                for key in (
                    "final_hallucination_rate",
                    "profile_resource_difficulty_adaptation_accuracy",
                    "strict_closed_loop_coverage",
                    "effective_automatic_adaptation_rate",
                    "hallucination_interception_rate",
                    "native_teaching_adaptation_mismatch_rate",
                )
            },
        },
    }
    facts = [{"content_unit_id": "CU-1", "case_id": "E2E-001", "content_text": "事实"}]
    nodes = [{"first_gen_transaction_id": "TX-1", "case_id": "E2E-001"}]
    cells = [{"coverage_cell_id": "CC-1", "seed_id": "seed_A", "case_id": "E2E-001"}]
    path = export_v4_workbook(tmp_path / "audit.xlsx", report, facts, nodes, cells)

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "00_指标汇总",
        "01_TRACE字段",
        "02_事实内容单元",
        "03_适配节点",
        "04_覆盖30格",
    ]
    for sheet_name, values in {
        "02_事实内容单元": ("SUPPORTED", "SUPPORTED", None),
        "03_适配节点": (0, 0, None),
        "04_覆盖30格": (1, 1, None),
    }.items():
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[3]]
        suffix = 5 if sheet_name == "02_事实内容单元" else 3
        for offset, value in enumerate(values):
            sheet.cell(4, len(headers) - suffix + 1 + offset, value)
    workbook.save(path)

    human = read_human_review_workbook(path)
    assert human["fact_units"][0]["reviewer_a"] == "SUPPORTED"
    assert human["adaptation_transactions"][0]["reviewer_a_mismatch"] == 0
    assert human["coverage_cells"][0]["reviewer_a_pass"] == 1
