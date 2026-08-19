"""Five-sheet V4 audit workbook export and double-review import."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Mapping, Sequence


METRIC_LABELS = {
    "final_hallucination_rate": "最终发布幻觉率",
    "profile_resource_difficulty_adaptation_accuracy": "画像—资源难度适配准确率",
    "strict_closed_loop_coverage": "主域严格闭环覆盖率",
    "effective_automatic_adaptation_rate": "动态有效调整率（诊断项）",
    "hallucination_interception_rate": "幻觉拦截率",
    "native_teaching_adaptation_mismatch_rate": "原生教学适配失配率",
}

TRACE_FIELD_ROWS = [
    ("公共标识", "run_id", "正式运行批次ID", "全部指标"),
    ("公共标识", "seed_id", "seed_A / seed_B", "两轮隔离"),
    ("公共标识", "case_id", "E2E案例ID", "全部指标"),
    ("公共标识", "session_id", "真实会话ID", "会话归属"),
    ("公共标识", "trace_id", "TRACE ID", "追溯"),
    ("公共标识", "msg_id", "消息ID", "追溯"),
    ("公共标识", "step", "事件序号", "排序"),
    ("公共标识", "timestamp", "ISO-8601时间", "排序"),
    ("公共标识", "agent", "执行Agent", "协同证据"),
    ("公共标识", "role", "消息角色", "协同证据"),
    ("载荷", "payload.type", "消息类型", "事实/适配/覆盖"),
    ("载荷", "payload.content.event", "业务事件", "适配节点"),
    ("路由", "route_mode", "production", "正式准入"),
    ("路由", "selected_knowledge_point", "确定性路由知识点", "适配/覆盖"),
    ("路由", "selected_difficulty", "初始难度", "适配"),
    ("路由", "knowledge_point_plan", "带证据的培养计划", "路由可追溯"),
    ("审核", "verdict.decision", "approve/reject", "发布/拦截"),
    ("审核", "verdict.rule_hits", "R-01～R-06", "幻觉/失配"),
    ("证据", "claims", "内容事实单元", "幻觉率"),
    ("证据", "evidence", "与claim绑定的证据", "幻觉/覆盖"),
    ("适配", "difficulty_action", "keep/step_up/step_down/deferred", "适配节点"),
    ("适配", "learner_follow_up_assessed", "学员回答判定", "适配节点"),
    ("复现", "code_version", "Git提交", "复现"),
    ("复现", "model_config", "模型与温度", "复现"),
]


def _json_cell(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A4" if worksheet.title != "00_指标汇总" else "A5"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for cell in worksheet[1]:
            cell.font = Font(size=15, bold=True, color="17365D")
        header_row = 4 if worksheet.title == "00_指标汇总" else 3
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in worksheet.columns:
            letter = column[0].column_letter
            width = min(max(12, max(len(str(cell.value or "")) for cell in column) + 2), 42)
            worksheet.column_dimensions[letter].width = width


def export_v4_workbook(
    path: Path,
    report: Mapping[str, Any],
    facts: Sequence[Mapping[str, Any]],
    nodes: Sequence[Mapping[str, Any]],
    cells: Sequence[Mapping[str, Any]],
) -> Path:
    from openpyxl import Workbook

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "00_指标汇总"
    summary.append(["TRACE V4｜AUTO_PRELIMINARY 与人工终审分栏"])
    summary.append([
        "自动初算不是最终成绩。人工终审必须完成双人复核；分歧必须仲裁。"
    ])
    summary.append([])
    summary.append([
        "指标", "自动分子", "自动分母", "自动百分比", "最终分子", "最终分母", "最终百分比", "状态"
    ])
    combined = report["combined"]
    automatic = report.get("automatic_preliminary") or combined
    for key, label in METRIC_LABELS.items():
        auto_metric = automatic["metrics"][key]
        final_metric = combined["metrics"][key]
        summary.append([
            label,
            auto_metric["numerator"],
            auto_metric["denominator"],
            None if auto_metric.get("denominator_zero") else auto_metric["percentage"] / 100,
            final_metric["numerator"] if report["mode"] == "FINAL_HUMAN_REVIEWED" else None,
            final_metric["denominator"] if report["mode"] == "FINAL_HUMAN_REVIEWED" else None,
            (
                final_metric["percentage"] / 100
                if report["mode"] == "FINAL_HUMAN_REVIEWED" and not final_metric.get("denominator_zero")
                else None
            ),
            combined["label_status"],
        ])
    for row in range(5, 11):
        summary.cell(row, 4).number_format = "0.0000%"
        summary.cell(row, 7).number_format = "0.0000%"

    trace = workbook.create_sheet("01_TRACE字段")
    trace.append(["生产TRACE字段规范与V4复算映射"])
    trace.append([])
    trace.append(["字段组", "JSON字段路径", "中文名称/允许值", "指标用途"])
    for row in TRACE_FIELD_ROWS:
        trace.append(list(row))

    fact_sheet = workbook.create_sheet("02_事实内容单元")
    fact_sheet.append(["事实内容单元｜自动标签保留，人工双审字段待填写"])
    fact_sheet.append([])
    fact_keys = list(facts[0].keys()) if facts else []
    fact_human = ["reviewer_a", "reviewer_b", "adjudicator_label", "human_reason", "human_rule_hits"]
    fact_sheet.append(fact_keys + fact_human)
    for row in facts:
        fact_sheet.append([_json_cell(row.get(key)) for key in fact_keys] + [None] * len(fact_human))

    node_sheet = workbook.create_sheet("03_适配节点")
    node_sheet.append(["适配节点｜金标仅在离线复算阶段合并"])
    node_sheet.append([])
    node_keys = list(nodes[0].keys()) if nodes else []
    node_human = ["reviewer_a_mismatch", "reviewer_b_mismatch", "adjudicator_mismatch"]
    node_sheet.append(node_keys + node_human)
    for row in nodes:
        node_sheet.append([_json_cell(row.get(key)) for key in node_keys] + [None] * len(node_human))

    cell_sheet = workbook.create_sheet("04_覆盖30格")
    cell_sheet.append(["10知识点×3难度｜每个Seed独立保留30格"])
    cell_sheet.append([])
    cell_keys = list(cells[0].keys()) if cells else []
    cell_human = ["reviewer_a_pass", "reviewer_b_pass", "adjudicator_pass"]
    cell_sheet.append(cell_keys + cell_human)
    for row in cells:
        cell_sheet.append([_json_cell(row.get(key)) for key in cell_keys] + [None] * len(cell_human))

    _style_workbook(workbook)
    temporary = target.with_name(f".{target.name}.tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(target)
    return target


def read_human_review_workbook(path: Path) -> dict[str, Any]:
    """Read reviewer fields only; automatic columns are deliberately ignored."""

    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=False)

    def records(sheet_name: str) -> list[dict[str, Any]]:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[3]]
        result: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=4, values_only=True):
            if not any(value is not None for value in values):
                continue
            result.append(dict(zip(headers, values)))
        return result

    facts = []
    for row in records("02_事实内容单元"):
        facts.append({
            "content_unit_id": row.get("content_unit_id"),
            "reviewer_a": row.get("reviewer_a"),
            "reviewer_b": row.get("reviewer_b"),
            "adjudicator": row.get("adjudicator_label"),
            "reason": row.get("human_reason") or "",
            "rule_hits": [
                item.strip()
                for item in str(row.get("human_rule_hits") or "").split(",")
                if item.strip()
            ],
        })

    transactions: dict[str, dict[str, Any]] = {}
    for row in records("03_适配节点"):
        tx_id = str(row.get("first_gen_transaction_id") or "")
        if not tx_id:
            continue
        current = transactions.setdefault(
            tx_id,
            {
                "first_gen_transaction_id": tx_id,
                "case_id": row.get("case_id"),
                "reviewer_a_mismatch": None,
                "reviewer_b_mismatch": None,
                "adjudicator_mismatch": None,
            },
        )
        for key in ("reviewer_a_mismatch", "reviewer_b_mismatch", "adjudicator_mismatch"):
            if row.get(key) is not None:
                current[key] = row.get(key)

    coverage = []
    for row in records("04_覆盖30格"):
        coverage.append({
            "coverage_cell_id": row.get("coverage_cell_id"),
            "seed_id": row.get("seed_id"),
            "case_id": row.get("case_id"),
            "reviewer_a_pass": row.get("reviewer_a_pass"),
            "reviewer_b_pass": row.get("reviewer_b_pass"),
            "adjudicator_pass": row.get("adjudicator_pass"),
        })
    return {
        "fact_units": facts,
        "adaptation_transactions": list(transactions.values()),
        "coverage_cells": coverage,
    }
